"""Write Agent C's review as hiring_review.xlsx, one sheet per section."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

# Mandatory requirements and core responsibilities count three times as much as preferred ones.
WEIGHTS = {"Mandatory": 3, "Core responsibility": 3, "Preferred": 1}
SCORES = {"STRONG": 1.0, "MODERATE": 0.6, "WEAK": 0.3, "MISSING": 0.0}
FILLS = {"STRONG": "C6EFCE", "MODERATE": "FFEB9C", "WEAK": "F8CBAD", "MISSING": "FFC7CE"}
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
WRAP = Alignment(wrap_text=True, vertical="top")


def weighted_coverage(requirements) -> float:
    total = sum(WEIGHTS[r["importance"]] for r in requirements)
    met = sum(WEIGHTS[r["importance"]] * SCORES[r["evidence_level"]] for r in requirements)
    return met / total if total else 0.0


def add_sheet(book, title, headers, rows, widths):
    sheet = book.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font, cell.fill = Font(bold=True), HEADER_FILL
    for column, width in zip("ABCDEFGH", widths):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    sheet.freeze_panes = "A2"
    return sheet


def write_review(review: dict, path: Path):
    book = Workbook()
    book.remove(book.active)
    requirements = review["requirements"]
    decision = review["interview_recommendation"]
    counts = {level: sum(r["evidence_level"] == level for r in requirements) for level in SCORES}

    summary = add_sheet(book, "Summary", ["Item", "Assessment"], [
        ["Weighted requirement coverage", f"{weighted_coverage(requirements):.0%}"],
        ["Evidence counts", ", ".join(f"{level} {n}" for level, n in counts.items())],
        ["Interview recommendation", f"{decision['decision']}: {decision['reason']}"],
        ["Overall assessment", review["overall_assessment"]],
        ["CV and cover letter together", review["combined_assessment"]],
        ["Final summary", review["final_summary"]],
        ["Coverage method", "Weights: Mandatory 3, Core responsibility 3, Preferred 1. "
                            "Scores: STRONG 1, MODERATE 0.6, WEAK 0.3, MISSING 0."],
    ], [32, 110])
    for cell in summary["A"]:
        cell.font = Font(bold=True)

    add_sheet(book, "Job analysis", ["Category", "Item"],
              [[a["category"], a["item"]] for a in review["job_analysis"]], [28, 100])

    sheet = add_sheet(book, "Requirements", ["Requirement", "Importance", "Evidence level", "Evidence", "Source", "Gap type", "Comment"],
                      [[r["requirement"], r["importance"], r["evidence_level"], r["evidence"], r["source"], r["gap_type"], r["comment"]]
                       for r in requirements], [40, 18, 15, 50, 13, 17, 50])
    for cell in sheet["C"][1:]:
        cell.fill = PatternFill("solid", fgColor=FILLS[cell.value])

    for title, key in (("CV assessment", "cv_assessment"), ("Cover letter assessment", "cover_letter_assessment")):
        add_sheet(book, title, ["Criterion", "Rating (1-5)", "Justification"],
                  [[a["criterion"], a["rating"], a["justification"]] for a in review[key]], [32, 13, 100])

    add_sheet(book, "ATS keywords", ["Keyword", "Status", "Note"],
              [[k["keyword"], k["status"], k["note"]] for k in review["ats_keywords"]], [30, 30, 80])
    add_sheet(book, "Strongest matches", ["Match"], [[m] for m in review["strongest_matches"]], [120])
    add_sheet(book, "Gaps", ["Gap", "Type", "Impact"],
              [[g["gap"], g["gap_type"], g["impact"]] for g in review["important_gaps"]], [50, 18, 70])
    add_sheet(book, "Recommendations", ["Document", "Recommendation", "Reason"],
              [["CV", r["recommendation"], r["reason"]] for r in review["cv_recommendations"]]
              + [["Cover letter", r["recommendation"], r["reason"]] for r in review["cover_letter_recommendations"]],
              [14, 70, 60])
    add_sheet(book, "Interview questions", ["Question", "Reason"],
              [[q["question"], q["reason"]] for q in review["interview_questions"]], [70, 60])
    book.save(path)
